"""管理端用户管理单元测试:分页筛选/详情统计/封禁链路/批量发券/风控/导出。

封禁链路是安全核心 —— 封禁必须同时做到:禁登录、踢下线、禁下单,三者缺一不可。
"""
from datetime import datetime, timedelta

import httpx
import pytest
from sqlalchemy import func, select

from app.core.exceptions import BizException, LoginFailedException
from app.main import app
from app.models import Coupon, Orders, User, UserCoupon, UserLoginLog
from app.schemas.business import OrdersSubmitIn
from app.services import order_service, risk_service, user_admin_service, user_service
from app.utils.password import hash_password

# ===== 辅助 =====


async def _stub_captcha(monkeypatch):
    """打桩验证码(通过)"""
    async def fake_verify(uuid, code):
        pass
    monkeypatch.setattr(user_service, "verify_captcha", fake_verify)


async def _user(db, username="zhangsan", phone="13800138000", status=1, password="abc12345"):
    u = User(username=username, phone=phone, password=hash_password(password), status=status)
    db.add(u)
    await db.commit()
    await db.refresh(u)
    return u


async def _coupon(db, stock=10, status=1, valid_days=7):
    now = datetime.now()
    c = Coupon(name="测试券", type=1, amount=10, min_amount=0, total=stock, stock=stock,
               per_user_limit=1, valid_days=valid_days,
               start_time=now - timedelta(days=1), end_time=now + timedelta(days=7),
               status=status)
    db.add(c)
    await db.commit()
    await db.refresh(c)
    return c


async def _log(db, user_id, ip, status=1, days_ago=0):
    db.add(UserLoginLog(user_id=user_id, login_type="account", status=status, ip=ip,
                        create_time=datetime.now() - timedelta(days=days_ago)))
    await db.commit()


def _order(user_id, number, status, amount):
    """构造订单(order_time 是 NOT NULL,必须给)"""
    return Orders(user_id=user_id, address_book_id=1, number=number, status=status,
                  amount=amount, order_time=datetime.now())


@pytest.fixture
async def admin_api(db):
    """带管理端鉴权的测试客户端:覆盖鉴权依赖(返回固定员工id)+ 注入测试 db 会话"""
    from app.core.database import get_db
    from app.core.security import get_current_admin

    async def _fake_admin():
        return 1

    async def _fake_db():
        yield db

    app.dependency_overrides[get_current_admin] = _fake_admin
    app.dependency_overrides[get_db] = _fake_db
    yield
    app.dependency_overrides.clear()


async def _call(method, path, **kw):
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as c:
        return await c.request(method, path, **kw)


# ===== 分页与详情 =====


async def test_page_filters(db):
    """分页筛选:用户名/手机号/状态"""
    await _user(db, "alice", "13800000001")
    await _user(db, "bob", "13800000002", status=0)

    total, rows = await user_admin_service.page_query(db)
    assert total == 2
    assert {r.username for r in rows} == {"alice", "bob"}

    total, rows = await user_admin_service.page_query(db, username="ali")
    assert total == 1 and rows[0].username == "alice"

    total, rows = await user_admin_service.page_query(db, phone="0002")
    assert total == 1 and rows[0].username == "bob"

    total, rows = await user_admin_service.page_query(db, status=0)
    assert total == 1 and rows[0].username == "bob"


async def test_page_filters_by_spend(db):
    """分群筛选:按累计消费额(只算已完成订单 status=5)"""
    u1 = await _user(db, "big", "13800000001")
    u2 = await _user(db, "small", "13800000002")
    db.add(_order(u1.id, "N1", 5, 500))
    db.add(_order(u1.id, "N2", 6, 999))  # 已取消不计
    db.add(_order(u2.id, "N3", 5, 10))
    await db.commit()

    total, rows = await user_admin_service.page_query(db, min_amount=100)
    assert total == 1 and rows[0].username == "big"

    total, rows = await user_admin_service.page_query(db, max_amount=100)
    assert total == 1 and rows[0].username == "small"


async def test_detail_stats(db):
    """详情统计:订单数(全部) / 已完成数 / 累计消费(只算 status=5)"""
    u = await _user(db)
    db.add(_order(u.id, "N1", 5, 100))
    db.add(_order(u.id, "N2", 5, 50))
    db.add(_order(u.id, "N3", 6, 999))
    await db.commit()

    d = await user_admin_service.get_detail(db, u.id)
    assert d["orderCount"] == 3        # 全部订单
    assert d["validOrderCount"] == 2   # 已完成
    assert float(d["totalSpend"]) == 150.0  # 已取消的 999 不计入
    assert "password" not in d
    assert d["status"] == 1


async def test_detail_not_found(db):
    with pytest.raises(BizException) as exc:
        await user_admin_service.get_detail(db, 99999)
    assert "不存在" in str(exc.value)


async def test_page_api_does_not_leak_password(db, admin_api):
    """回归防护:分页接口响应体里不能出现 password(HASH 泄露)

    员工分页接口曾因漏加 exclude 把 bcrypt 哈希返回给前端 —— 这里从 HTTP 响应层面守死。
    """
    await _user(db)
    r = await _call("GET", "/admin/user/page")
    assert r.status_code == 200
    body = r.json()
    assert body["code"] == 1
    record = body["data"]["records"][0]
    assert "password" not in record
    # 确认确实返回了用户(而不是空列表导致断言假通过)
    assert record["username"] == "zhangsan"


# ===== 封禁链路(安全核心)=====


async def test_ban_rejects_login(db, monkeypatch):
    """封禁后:密码正确也登录失败"""
    await _stub_captcha(monkeypatch)
    u = await _user(db, status=0)
    with pytest.raises(LoginFailedException) as exc:
        await user_service.login(db, "zhangsan", "abc12345")
    assert "封禁" in str(exc.value)


async def test_ban_clears_session(db, monkeypatch):
    """封禁即踢下线:清 session:user:{id}(否则 JWT 还能用满 2 小时)"""
    u = await _user(db)
    deleted = []

    async def _record(key):
        deleted.append(key)

    monkeypatch.setattr("app.core.redis.redis_delete", _record)
    await user_admin_service.change_status(db, 1, u.id, 0, "恶意刷单")

    assert f"session:user:{u.id}" in deleted
    await db.refresh(u)
    assert u.status == 0
    assert u.ban_reason == "恶意刷单"
    assert u.update_user == 1  # 记录操作人(审计)


async def test_unban_allows_login(db, monkeypatch):
    """解封后可正常登录,且封禁原因被清空"""
    await _stub_captcha(monkeypatch)
    u = await _user(db, status=0)
    await user_admin_service.change_status(db, 1, u.id, 0, "误封")
    await user_admin_service.change_status(db, 1, u.id, 1)

    await db.refresh(u)
    assert u.status == 1
    assert u.ban_reason is None

    data = await user_service.login(db, "zhangsan", "abc12345")
    assert data["token"]


async def test_login_writes_last_login(db, monkeypatch):
    """登录成功回写 last_login_time / last_login_ip(详情展示与活跃度筛选用)"""
    await _stub_captcha(monkeypatch)
    u = await _user(db)
    await user_service.login(db, "zhangsan", "abc12345", ip="1.2.3.4")
    await db.refresh(u)
    assert u.last_login_time is not None
    assert u.last_login_ip == "1.2.3.4"


async def test_banned_user_cannot_submit_order(db):
    """下单纵深防御:Redis 不可用时会话校验会降级放行,封禁账号仍必须被拦在下单之外"""
    u = await _user(db, status=0)
    with pytest.raises(BizException) as exc:
        await order_service.submit(db, u.id, OrdersSubmitIn(address_book_id=1, amount=100))
    assert "封禁" in str(exc.value)


async def test_ban_not_found(db):
    with pytest.raises(BizException) as exc:
        await user_admin_service.change_status(db, 1, 99999, 0)
    assert "不存在" in str(exc.value)


async def test_ban_invalid_status(db):
    u = await _user(db)
    with pytest.raises(BizException) as exc:
        await user_admin_service.change_status(db, 1, u.id, 2)
    assert "状态参数" in str(exc.value)


# ===== 批量发券 =====


async def test_grant_coupon_success(db):
    """发券成功:库存扣减 + 券入库"""
    u1 = await _user(db, "u1", "13800000001")
    u2 = await _user(db, "u2", "13800000002")
    c = await _coupon(db, stock=10)

    r = await user_admin_service.grant_coupon(db, 1, c.id, [u1.id, u2.id])
    assert r["granted"] == 2
    assert r["stockLeft"] == 8

    n = await db.scalar(select(func.count(UserCoupon.id)).where(UserCoupon.coupon_id == c.id))
    assert n == 2
    row = (await db.execute(select(UserCoupon).where(UserCoupon.user_id == u1.id))).scalar_one()
    assert row.expire_time is not None  # 过期时间=领取时间+有效天数


async def test_grant_coupon_skips_claimed(db):
    """已领过的用户被跳过,不重复发(唯一约束的业务层前置)"""
    u1 = await _user(db, "u1", "13800000001")
    u2 = await _user(db, "u2", "13800000002")
    c = await _coupon(db, stock=10)
    db.add(UserCoupon(user_id=u1.id, coupon_id=c.id, status=0))
    await db.commit()

    r = await user_admin_service.grant_coupon(db, 1, c.id, [u1.id, u2.id])
    assert r["granted"] == 1 and r["skipped"] == 1


async def test_grant_coupon_dedup_input(db):
    """入参重复 id 去重(否则会多发)"""
    u = await _user(db)
    c = await _coupon(db, stock=10)
    r = await user_admin_service.grant_coupon(db, 1, c.id, [u.id, u.id, u.id])
    assert r["granted"] == 1


async def test_grant_coupon_insufficient_stock(db):
    """库存不足:整批拒绝,一张不发(不做部分发放,避免运营对不上账)"""
    u1 = await _user(db, "u1", "13800000001")
    u2 = await _user(db, "u2", "13800000002")
    u3 = await _user(db, "u3", "13800000003")
    c = await _coupon(db, stock=2)
    # 先取出 id:库存不足时服务会 rollback,rollback 会让 ORM 对象过期,
    # 之后再访问 c.id 会触发懒加载 IO(MissingGreenlet)
    coupon_id = c.id

    with pytest.raises(BizException) as exc:
        await user_admin_service.grant_coupon(db, 1, coupon_id, [u1.id, u2.id, u3.id])
    assert "库存不足" in str(exc.value)

    n = await db.scalar(select(func.count(UserCoupon.id)))
    assert n == 0  # 一张都没发出去
    assert (await db.scalar(select(Coupon.stock).where(Coupon.id == coupon_id))) == 2  # 库存未动


async def test_grant_coupon_disabled_or_expired(db):
    """停用 / 已过期的券不能发放"""
    u = await _user(db)
    disabled = await _coupon(db, status=0)
    with pytest.raises(BizException) as exc:
        await user_admin_service.grant_coupon(db, 1, disabled.id, [u.id])
    assert "停用" in str(exc.value)

    now = datetime.now()
    expired = Coupon(name="过期券", type=1, amount=10, min_amount=0, total=5, stock=5,
                     per_user_limit=1, valid_days=7, status=1,
                     start_time=now - timedelta(days=10), end_time=now - timedelta(days=1))
    db.add(expired)
    await db.commit()
    with pytest.raises(BizException) as exc:
        await user_admin_service.grant_coupon(db, 1, expired.id, [u.id])
    assert "可领时间" in str(exc.value)


# ===== 风控 =====


async def test_risk_multi_ip(db):
    """信号1:同一账号从多个不同 IP 登录"""
    u = await _user(db)
    for ip in ("1.1.1.1", "2.2.2.2", "3.3.3.3"):
        await _log(db, u.id, ip)

    risks = await risk_service.risk_users(db)
    assert len(risks) == 1
    assert risks[0]["userId"] == u.id
    assert risks[0]["distinctIpCount"] == 3
    assert any("不同 IP" in s for s in risks[0]["signals"])
    assert risks[0]["username"] == "zhangsan"


async def test_risk_fail_rate(db):
    """信号3:登录失败率过高(样本不足时不误报)"""
    u = await _user(db)
    for _ in range(4):
        await _log(db, u.id, "1.1.1.1", status=0)
    await _log(db, u.id, "1.1.1.1", status=1)

    risks = await risk_service.risk_users(db)
    assert len(risks) == 1
    assert any("失败率" in s for s in risks[0]["signals"])

    # 样本不足(仅 2 次失败)不应触发
    u2 = await _user(db, "lisi", "13800000002")
    for _ in range(2):
        await _log(db, u2.id, "9.9.9.9", status=0)
    risks = await risk_service.risk_users(db)
    assert all(r["userId"] != u2.id for r in risks)


async def test_risk_shared_ip(db):
    """信号2:同一 IP 登录多个账号(撞库特征)"""
    users = [await _user(db, f"u{i}", f"138000000{i:02d}") for i in range(5)]
    for u in users:
        await _log(db, u.id, "6.6.6.6")

    risks = await risk_service.risk_users(db)
    ids = {r["userId"] for r in risks}
    assert ids == {u.id for u in users}
    assert all(any("共用" in s for s in r["signals"]) for r in risks)
    # 该 IP 被标记为 shared
    assert risks[0]["ips"][0]["shared"] is True


async def test_risk_normal_user_not_flagged(db):
    """正常用户(单一 IP、无失败)不进风险列表"""
    u = await _user(db)
    for _ in range(5):
        await _log(db, u.id, "1.1.1.1")
    assert await risk_service.risk_users(db) == []


async def test_risk_level_two_signals_is_high(db):
    """命中 ≥2 个信号 → RISK(高危);1 个 → WATCH(关注)"""
    u = await _user(db)
    # 多 IP(3个)+ 高失败率 = 2 个信号
    for ip in ("1.1.1.1", "2.2.2.2", "3.3.3.3"):
        for _ in range(4):
            await _log(db, u.id, ip, status=0)
    risks = await risk_service.risk_users(db)
    assert risks[0]["riskLevel"] == "RISK"


# ===== 导出 =====


async def test_export_excel(db):
    """导出 xlsx:含表头,中文不乱码,手机号按文本存"""
    from openpyxl import load_workbook
    from io import BytesIO

    await _user(db, "张三", "13800138000")
    content, filename = await user_admin_service.build_export_excel(db)

    assert filename.endswith(".xlsx")
    assert len(content) > 1000  # 不是空文件
    ws = load_workbook(BytesIO(content)).active
    header = [c.value for c in ws[1]]
    assert header[0] == "ID" and "用户名" in header and "累计消费(元)" in header
    row = [c.value for c in ws[2]]
    assert row[1] == "张三"
    assert row[2] == "13800138000"


async def test_export_api_route_not_shadowed(db, admin_api):
    """回归防护:/admin/user/export 不能被 /{user_id} 抢先匹配(否则会被解析成 int 报 422)"""
    await _user(db)
    r = await _call("GET", "/admin/user/export")
    assert r.status_code == 200, f"路由被遮蔽,状态码 {r.status_code}"
    assert "spreadsheetml" in r.headers.get("content-type", "")


async def test_risk_api(db, admin_api):
    """风险列表接口可用"""
    u = await _user(db)
    for ip in ("1.1.1.1", "2.2.2.2", "3.3.3.3"):
        await _log(db, u.id, ip)
    r = await _call("GET", "/admin/user/risk")
    assert r.status_code == 200
    data = r.json()["data"]
    assert len(data) == 1 and data[0]["riskLevel"] in ("WATCH", "RISK")
